#!/usr/bin/env python3
"""Na mídia: planilha data/midia.xlsx <-> index.html.

    python3 scripts/midia.py check            # valida a planilha, não altera nada
    python3 scripts/midia.py check --links    # também abre cada URL publicada e atualiza "Checado em" na planilha
    python3 scripts/midia.py build            # regenera a página /midia e a seção Na mídia da home
    python3 scripts/midia.py export           # recria a planilha a partir do index.html (use --force para sobrescrever)

A planilha é a fonte da verdade. `build` só reescreve os blocos entre <!-- MIDIA:X --> e
<!-- /MIDIA:X --> do index.html: FAIXA (wordmarks) e DESTAQUES na home; TOTAL, FILTROS e
LISTA na página /midia.

O portfolio.py e o midia.py só conhecem o index.html. Ordem certa:

    python3 scripts/portfolio.py build
    python3 scripts/midia.py build
    python3 scripts/pages.py build
"""
import argparse
import concurrent.futures
import datetime
import html
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request

try:
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("Instale o openpyxl: pip3 install openpyxl")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HTML_PATH = os.path.join(ROOT, "index.html")
XLSX_PATH = os.path.join(ROOT, "data", "midia.xlsx")
SHEET = "Matérias"
SHEET_V = "Veículos"

# (código, título do bloco na página, linha de contexto). A ordem aqui é a ordem dos
# blocos na página e dos botões de filtro.
CATEGORIAS = [
    ("perfil", "Perfil", "Quem eu sou. Perfis, entrevistas biográficas e prêmios."),
    ("tese", "Tese", "Como eu penso. Entrevistas e análises sobre mercado, venture capital e IA."),
    ("movimento", "Movimento", "O que construí. Da venda do Buscapé à Redpoint, à Headline e às captações dos fundos."),
    ("portfolio", "Portfólio", "O fundo operando. Aportes, rodadas e saídas de empresas investidas, com o meu nome na matéria."),
]
ASSINADOS = ("Textos assinados", "O que escrevi para os veículos. Os demais textos estão no ")
SUBSTACK = "https://substack.com/@romerorodrigues"
CAT_CODES = [c[0] for c in CATEGORIAS]
CAT_LABEL = {c[0]: c[1] for c in CATEGORIAS}
FORMATOS = {"texto": "Texto", "video": "Vídeo", "audio": "Áudio"}
FORMATO_TAG = {"video": "vídeo", "audio": "áudio"}
IDIOMAS = {"pt": ("pt-BR", "Em português"), "en": ("en", "Em inglês")}
PAISES = {"BR", "INT"}
MESES = {"pt": "jan fev mar abr mai jun jul ago set out nov dez".split(),
         "en": "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()}
TAG_PAYWALL = "assinantes"
TAG_ASSINADO = "por Romero"
DESTAQUES = 4  # cards da seção Na mídia na home
TRACKING = re.compile(r"^(utm_\w+|fbclid|gclid|dclid|mc_cid|mc_eid|igshid|_ga|ref_src|cmpid)$", re.I)

# (chave, cabeçalho, largura, comentário)
COLUMNS = [
    ("publicar", "Publicar", 9, "x = aparece no site. Vazio = fica só na planilha."),
    ("destaque", "Destaque", 10, "1 a 4 = posição na home. Vazio = fora da home. Precisa haver exatamente quatro."),
    ("data", "Data", 12, "AAAA-MM-DD. Ordenação e exibição saem daqui. Sem o dia, use AAAA-MM (ou AAAA)."),
    ("veiculo", "Veículo", 20, "Precisa bater com a aba Veículos."),
    ("titulo", "Título", 60, "Como o veículo publicou, sem reescrever."),
    ("url", "URL", 50, "Link direto, sem parâmetro de rastreio (utm_, fbclid...)."),
    ("categoria", "Categoria", 12, "perfil, tese, movimento ou portfolio: o que a peça prova sobre você."),
    ("formato", "Formato", 10, "texto, video ou audio."),
    ("idioma", "Idioma", 8, "pt ou en."),
    ("assinado", "Assinado", 10, "x = escrito por você. Ganha 'por Romero' e vai para Textos assinados."),
    ("empresa", "Empresa", 18, "Só na categoria portfolio. Nome como está no portfolio.xlsx."),
    ("paywall", "Paywall", 9, "x = marca 'assinantes' no item."),
    ("subjectof", "subjectOf", 10, "x = entra no subjectOf da Person no JSON-LD de todas as páginas. No máximo 5, nunca texto assinado."),
    ("nota", "Nota", 30, "Uma linha. Só usada nos destaques da home (e precisa de tradução em data/en.json)."),
    ("checado", "Checado em", 12, "Data da última vez que a URL respondeu 200. Atualizada por: midia.py check --links"),
    ("obs", "Observações", 40, "Interna, não vai para o site."),
]
KEYS = [c[0] for c in COLUMNS]
POSICOES = [str(i) for i in range(1, DESTAQUES + 1)]
HEADERS = {c[0]: c[1] for c in COLUMNS}
FLAGS = ["publicar", "assinado", "paywall", "subjectof"]

COLUMNS_V = [
    ("veiculo", "Veículo", 22, "Chave, igual à coluna Veículo da aba Matérias."),
    ("wordmark", "Wordmark", 22, "Como escrever o nome na faixa da home. O site mostra em caixa alta."),
    ("faixa", "Faixa", 8, "Número = posição na faixa da home. Vazio = fora da faixa."),
    ("site", "Site", 34, "Home do veículo. Vai para o publisher do JSON-LD."),
    ("pais", "País", 8, "BR ou INT."),
]
KEYS_V = [c[0] for c in COLUMNS_V]
HEADERS_V = {c[0]: c[1] for c in COLUMNS_V}

BLOCKS = ("FAIXA", "DESTAQUES", "TOTAL", "FILTROS", "LISTA")


def block_re(name):
    return re.compile(r"(<!-- MIDIA:%s -->)(.*?)(<!-- /MIDIA:%s -->)" % (name, name), re.S)


def esc(s):
    return html.escape(s, quote=True)


# ---------------------------------------------------------------- datas

def parse_date(s):
    """'AAAA-MM-DD', 'AAAA-MM' ou 'AAAA' -> (ano, mês|None, dia|None), ou None se inválida."""
    m = re.fullmatch(r"(\d{4})(?:-(\d{2})(?:-(\d{2}))?)?", s)
    if not m:
        return None
    y, mo, d = (int(x) if x else None for x in m.groups())
    try:
        datetime.date(y, mo or 1, d or 1)
    except ValueError:
        return None
    return y, mo, d


def fmt_date(s, lang="pt"):
    """'2015-11-10' -> 'nov 2015' (ou 'Nov 2015' em inglês); '2023' -> '2023'."""
    y, mo, _ = parse_date(s)
    return f"{MESES[lang][mo - 1]} {y}" if mo else str(y)


def total_phrase(pub, lang="pt"):
    first = min(parse_date(r["data"])[0] for r in pub)
    return (f"{len(pub)} matérias, entrevistas e podcasts desde {first}" if lang == "pt" else
            f"{len(pub)} articles, interviews and podcasts since {first}")


# ---------------------------------------------------------------- planilha -> dados

def cell_str(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def marked(v):
    return cell_str(v).lower() in {"x", "sim", "s", "ok", "1", "yes", "y", "true", "verdadeiro", "✓", "✔"}


def sheet_rows(ws, keys, headers, first):
    header = [cell_str(c.value) for c in ws[1]]
    by_header = {h: i for i, h in enumerate(header)}
    missing = [headers[k] for k in keys if headers[k] not in by_header]
    if missing:
        sys.exit(f"Colunas faltando na aba {ws.title}: {', '.join(missing)}")
    for n, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        r = {k: raw[by_header[headers[k]]] if by_header[headers[k]] < len(raw) else None for k in keys}
        if cell_str(r[first]):
            yield n, r


def read_xlsx(path):
    """(matérias, {veículo: dados})."""
    wb = openpyxl.load_workbook(path, data_only=True)
    for name in (SHEET, SHEET_V):
        if name not in wb.sheetnames:
            sys.exit(f"Aba {name} não existe em {os.path.relpath(path, ROOT)}")
    rows = []
    for n, raw in sheet_rows(wb[SHEET], KEYS, HEADERS, "titulo"):
        row = {k: cell_str(raw[k]) for k in KEYS}
        for k in FLAGS:
            row[k] = marked(raw[k])
        for k in ("categoria", "formato", "idioma"):
            row[k] = row[k].lower()
        row["_line"] = n
        rows.append(row)
    vehicles = {}
    for n, raw in sheet_rows(wb[SHEET_V], KEYS_V, HEADERS_V, "veiculo"):
        v = {k: cell_str(raw[k]) for k in KEYS_V}
        v["_line"] = n
        if v["veiculo"] in vehicles:
            v["_dup"] = vehicles[v["veiculo"]]["_line"]
        vehicles[v["veiculo"]] = v
    return rows, vehicles


def num(s):
    try:
        return float(s)
    except ValueError:
        return None


def published(rows):
    """Matérias publicadas, da mais recente para a mais antiga."""
    pub = [r for r in rows if r["publicar"]]
    pub.sort(key=lambda r: r["titulo"].lower())
    pub.sort(key=lambda r: r["data"], reverse=True)
    return pub


def highlights(pub):
    return sorted((r for r in pub if r["destaque"]), key=lambda r: num(r["destaque"]) or 0)


def strip_items(vehicles):
    return sorted((v for v in vehicles.values() if v["faixa"]), key=lambda v: num(v["faixa"]) or 0)


def portfolio_names():
    """Nomes das empresas do portfolio.xlsx, para conferir a coluna Empresa. None se não der para ler."""
    try:
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            import portfolio
            return {r["nome"] for r in portfolio.read_xlsx(portfolio.XLSX_PATH)}
    except (Exception, SystemExit):
        return None


def validate(rows, vehicles):
    errors, warnings = [], []
    for v in vehicles.values():
        where = f"{SHEET_V}, linha {v['_line']} ({v['veiculo']})"
        if "_dup" in v:
            errors.append(f"{where}: veículo repetido (também na linha {v['_dup']})")
        if v["pais"] and v["pais"] not in PAISES:
            errors.append(f"{where}: País deve ser BR ou INT, não {v['pais']!r}")
        if v["site"] and not re.match(r"https?://", v["site"]):
            errors.append(f"{where}: Site deve começar com http:// ou https://")
        if v["faixa"]:
            if num(v["faixa"]) is None:
                errors.append(f"{where}: Faixa deve ser número")
            if not v["wordmark"]:
                errors.append(f"{where}: está na faixa sem Wordmark")
    pos = {}
    for v in strip_items(vehicles):
        if v["faixa"] in pos:
            errors.append(f"{SHEET_V}: {v['veiculo']} e {pos[v['faixa']]} na mesma posição da faixa")
        pos[v["faixa"]] = v["veiculo"]

    urls, titles, companies = {}, {}, None
    for r in rows:
        where = f"linha {r['_line']} ({r['titulo'][:48]})"
        u = r["url"]
        if u in urls:
            errors.append(f"{where}: URL repetida (também na linha {urls[u]})")
        urls[u] = r["_line"]
        if not re.match(r"https?://", u):
            errors.append(f"{where}: URL deve começar com http:// ou https://")
        else:
            tracking = [k for k, _ in urllib.parse.parse_qsl(urllib.parse.urlsplit(u).query) if TRACKING.match(k)]
            if tracking:
                errors.append(f"{where}: URL com parâmetro de rastreio: {', '.join(tracking)}")
        if r["categoria"] and r["categoria"] not in CAT_CODES:
            errors.append(f"{where}: Categoria {r['categoria']!r} fora de {', '.join(CAT_CODES)}")
        if r["formato"] and r["formato"] not in FORMATOS:
            errors.append(f"{where}: Formato {r['formato']!r} fora de {', '.join(FORMATOS)}")
        if r["idioma"] and r["idioma"] not in IDIOMAS:
            errors.append(f"{where}: Idioma {r['idioma']!r} fora de {', '.join(IDIOMAS)}")
        if r["data"] and not parse_date(r["data"]):
            errors.append(f"{where}: Data fora do formato AAAA-MM-DD: {r['data']!r}")
        if r["checado"] and not parse_date(r["checado"]):
            warnings.append(f"{where}: Checado em fora do formato AAAA-MM-DD: {r['checado']!r}")
        if r["veiculo"] and r["veiculo"] not in vehicles:
            errors.append(f"{where}: veículo {r['veiculo']!r} não está na aba {SHEET_V}")
        if not r["publicar"]:
            for k in ("destaque", "subjectof"):
                if r[k]:
                    errors.append(f"{where}: {HEADERS[k]} marcado numa matéria que não está publicada")
            continue
        for k in ("data", "veiculo", "url", "categoria", "formato", "idioma"):
            if not r[k]:
                errors.append(f"{where}: {HEADERS[k]} vazio")
        if r["titulo"] in titles:
            errors.append(f"{where}: título repetido (também na linha {titles[r['titulo']]})")
        titles[r["titulo"]] = r["_line"]
        if r["data"] and parse_date(r["data"]) and not parse_date(r["data"])[2]:
            warnings.append(f"{where}: data sem o dia ({r['data']}); a ordenação fica aproximada")
        if r["empresa"]:
            if r["categoria"] != "portfolio":
                warnings.append(f"{where}: Empresa preenchida fora da categoria portfolio")
            if companies is None:
                companies = portfolio_names() or set()
            if companies and r["empresa"] not in companies:
                warnings.append(f"{where}: empresa {r['empresa']!r} não está no portfolio.xlsx")
        if r["subjectof"] and r["assinado"]:
            errors.append(f"{where}: texto assinado não entra no subjectOf (é autoria, não assunto)")
        if r["destaque"] and r["destaque"] not in POSICOES:
            errors.append(f"{where}: Destaque deve ser de 1 a {DESTAQUES}")

    pub = published(rows)
    marks = sorted(r["destaque"] for r in pub if r["destaque"])
    if marks != POSICOES:
        errors.append(f"Destaque: precisa de exatamente {DESTAQUES} matérias publicadas, nas posições {', '.join(POSICOES)} (hoje: {', '.join(marks) or 'nenhuma'})")
    n_subject = sum(r["subjectof"] for r in pub)
    if n_subject > 5:
        errors.append(f"subjectOf: {n_subject} matérias marcadas, o máximo é 5")
    live = {r["veiculo"] for r in pub}
    for v in strip_items(vehicles):
        if v["veiculo"] not in live:
            warnings.append(f"Faixa: {v['veiculo']} está na faixa da home sem nenhuma matéria publicada")
    if not pub:
        errors.append("nenhuma matéria publicada")
    return errors, warnings


# ---------------------------------------------------------------- dados -> HTML

def lang_of(r):
    return IDIOMAS[r["idioma"]][0]


def tags(r):
    out = []
    if r["assinado"]:
        out.append(TAG_ASSINADO)
    if r["formato"] in FORMATO_TAG:
        out.append(FORMATO_TAG[r["formato"]])
    if r["paywall"]:
        out.append(TAG_PAYWALL)
    return out


def cats(r):
    return r["categoria"] + (" assinado" if r["assinado"] else "")


def render_item(r):
    name = " ".join(x for x in (r["veiculo"], r["titulo"], r["empresa"]) if x)
    tg = "".join(f"<span>{esc(t)}</span>" for t in tags(r))
    return (f'<li data-cats="{cats(r)}" data-name="{esc(name)}" data-date="{r["data"]}" data-ano="{parse_date(r["data"])[0]}" '
            f'data-formato="{r["formato"]}" data-idioma="{r["idioma"]}">'
            f'<a class="press-row" href="{esc(r["url"])}" target="_blank" rel="noopener" data-ga="midia_item" '
            f'data-veiculo="{esc(r["veiculo"])}" data-categoria="{r["categoria"]}">'
            f'<time datetime="{r["data"]}">{fmt_date(r["data"])}</time><span class="src">{esc(r["veiculo"])}</span>'
            f'<span class="ttl" lang="{lang_of(r)}">{esc(r["titulo"])}</span><span class="tags">{tg}</span></a></li>')


def render_group(code, title, lede, items, extra=""):
    lis = "".join(render_item(r) for r in items)
    return (f'<section class="press-group" data-group="{code}" aria-labelledby="midia-{code}">'
            f'<div class="press-head"><h2 id="midia-{code}">{esc(title)}</h2><p>{lede}{extra}</p></div>'
            f'<ol class="press">{lis}</ol></section>')


def render_list(pub):
    parts = []
    for code, title, lede in CATEGORIAS:
        items = [r for r in pub if r["categoria"] == code and not r["assinado"]]
        if items:
            parts.append(render_group(code, title, esc(lede), items))
    signed = [r for r in pub if r["assinado"]]
    if signed:
        link = f'<a class="link" href="{SUBSTACK}" target="_blank" rel="noopener">Substack</a>.'
        parts.append(render_group("assinado", ASSINADOS[0], esc(ASSINADOS[1]), signed, link))
    parts.append('<ol class="press press-flat" hidden></ol>')
    return "".join(parts)


def render_filters(pub):
    btns = [("all", "Tudo", len(pub))] + [(c, CAT_LABEL[c], sum(r["categoria"] == c for r in pub)) for c in CAT_CODES]
    buttons = "".join(
        f'<button type="button" data-filter="{c}" aria-pressed="{"true" if c == "all" else "false"}">{label}<sup>{n}</sup></button>'
        for c, label, n in btns if n or c == "all")
    search = '<label class="search"><span class="sr-only">Buscar veículo ou título</span><input type="search" placeholder="Buscar veículo ou título" autocomplete="off"></label>'

    def select(key, label, everyone, options):
        opts = "".join(f'<option value="{esc(v)}">{esc(t)}</option>' for v, t in options)
        return f'<select data-key="{key}" aria-label="{label}"><option value="">{everyone}</option>{opts}</select>'

    formatos = [(k, v) for k, v in FORMATOS.items() if any(r["formato"] == k for r in pub)]
    idiomas = [(k, v[1]) for k, v in IDIOMAS.items() if any(r["idioma"] == k for r in pub)]
    anos = sorted({str(parse_date(r["data"])[0]) for r in pub}, reverse=True)
    tools = ('<div class="press-tools">'
             + select("formato", "Formato", "Todos os formatos", formatos)
             + select("idioma", "Idioma", "Todos os idiomas", idiomas)
             + select("ano", "Ano", "Todos os anos", [(a, a) for a in anos])
             + '<div class="press-order" role="group" aria-label="Ordenar">'
               '<button type="button" data-order="cat" aria-pressed="true">Por categoria</button>'
               '<button type="button" data-order="date" aria-pressed="false">Por data</button></div></div>')
    return (f'<div class="filters" role="group" aria-label="Filtrar por categoria">{buttons}{search}</div>{tools}')


def render_strip(vehicles):
    names = "".join(f"<span>{esc(v['wordmark'])}</span>" for v in strip_items(vehicles))
    return (f'<a class="wordmarks" href="/midia" data-ga="midia_faixa"><span class="wordmarks-label">Cobertura em</span>'
            f'<span class="wordmarks-list">{names}</span></a>')


def render_highlights(pub):
    lis = []
    for r in highlights(pub):
        src = f'<span>{esc(r["veiculo"])}</span> · <time datetime="{r["data"]}">{fmt_date(r["data"])}</time>'
        if r["paywall"]:
            src += f" · <span>{TAG_PAYWALL}</span>"
        note = f'<small>{esc(r["nota"])}</small>' if r["nota"] else ""
        lis.append(f'<li><a class="row" href="{esc(r["url"])}" target="_blank" rel="noopener" data-ga="midia_destaque" '
                   f'data-veiculo="{esc(r["veiculo"])}" data-posicao="{r["destaque"]}">'
                   f'<span class="src">{src}</span><span class="ttl" lang="{lang_of(r)}">{esc(r["titulo"])}{note}</span>'
                   f'<span class="go">Ler</span></a></li>')
    return f'<ul class="rows">{"".join(lis)}</ul>'


def build(rows, vehicles, dry=False):
    pub = published(rows)
    content = {
        "FAIXA": render_strip(vehicles),
        "DESTAQUES": render_highlights(pub),
        "TOTAL": esc(total_phrase(pub)),
        "FILTROS": render_filters(pub),
        "LISTA": render_list(pub),
    }
    src = open(HTML_PATH, encoding="utf-8").read()
    out = src
    for name in BLOCKS:
        out, n = block_re(name).subn(lambda m: m.group(1) + content[name] + m.group(3), out)
        if n != 1:
            sys.exit(f"Não encontrei no index.html o bloco <!-- MIDIA:{name} --> ... <!-- /MIDIA:{name} --> ({n} ocorrências)")
    changed = out != src
    if changed and not dry:
        open(HTML_PATH, "w", encoding="utf-8").write(out)
    return pub, changed


# ---------------------------------------------------------------- dados para o pages.py

def en_textos(pub):
    """Traduções que dependem dos dados: datas dos itens e a frase do total."""
    out = {fmt_date(r["data"]): fmt_date(r["data"], "en") for r in pub}
    out[total_phrase(pub)] = total_phrase(pub, "en")
    return out


def keep_names(pub, vehicles):
    """Textos que não se traduzem: veículos, wordmarks e títulos, que ficam como o veículo publicou."""
    return {x for r in pub for x in (r["veiculo"], r["titulo"])} | {v["wordmark"] for v in vehicles.values() if v["wordmark"]}


def creative_work(r, vehicles, person):
    """Nó schema.org de uma matéria: NewsArticle (texto), PodcastEpisode (áudio) ou CreativeWork (vídeo).

    Vídeo não vira VideoObject: o Google pede que o vídeo marcado esteja na própria página,
    e aqui ele só é linkado."""
    kind = {"texto": "NewsArticle", "audio": "PodcastEpisode"}.get(r["formato"], "CreativeWork")
    node = {"@type": kind, ("headline" if kind == "NewsArticle" else "name"): r["titulo"], "url": r["url"],
            "datePublished": r["data"], "inLanguage": lang_of(r)}
    pub = {"@type": "Organization", "name": r["veiculo"]}
    site = vehicles.get(r["veiculo"], {}).get("site")
    if site:
        pub["url"] = site
    node["publisher"] = pub
    node["author" if r["assinado"] else "about"] = person
    return node


# ---------------------------------------------------------------- links

UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0 Safari/537.36"


def probe(url):
    """(status, detalhe). YouTube é conferido pelo oEmbed, que responde 404 para vídeo removido."""
    target = url
    if re.match(r"https?://(www\.)?youtube\.com/watch", url):
        target = "https://www.youtube.com/oembed?format=json&url=" + urllib.parse.quote(url, safe="")
    req = urllib.request.Request(target, headers={"User-Agent": UA, "Accept-Language": "pt-BR,pt;q=0.9,en;q=0.8"})
    try:
        with urllib.request.urlopen(req, timeout=25) as r:
            final = r.geturl()
            return r.status, (f"redireciona para {final}" if target == url and final.rstrip("/") != url.rstrip("/") else "")
    except urllib.error.HTTPError as e:
        return e.code, ""
    except Exception as e:  # rede, DNS, timeout
        return None, type(e).__name__


def check_links(path, rows, only_published=True):
    todo = [r for r in rows if r["publicar"] or not only_published]
    print(f"conferindo {len(todo)} links...")
    with concurrent.futures.ThreadPoolExecutor(12) as ex:
        results = list(ex.map(lambda r: probe(r["url"]), todo))
    today = datetime.date.today().isoformat()
    ok, errors = [], []
    for r, (status, detail) in zip(todo, results):
        where = f"linha {r['_line']} ({r['veiculo']}: {r['titulo'][:48]})"
        if status and 200 <= status < 300:
            ok.append(r["_line"])
            if detail:
                print(f"aviso: {where}: {detail}")
        elif status in (401, 403, 429, 999):
            print(f"aviso: {where}: {status}, o veículo bloqueia robôs. Confira no navegador.")
        elif status:
            errors.append(f"{where}: responde {status}")
        else:
            print(f"aviso: {where}: sem resposta ({detail})")
    lock = os.path.join(os.path.dirname(path), "~$" + os.path.basename(path))
    if ok:
        if os.path.exists(lock):
            print(f"aviso: {os.path.basename(path)} parece aberta no Excel; feche antes para não perder a atualização")
        wb = openpyxl.load_workbook(path)
        ws = wb[SHEET]
        col = [cell_str(c.value) for c in ws[1]].index(HEADERS["checado"]) + 1
        for line in ok:
            ws.cell(row=line, column=col, value=today)
        wb.save(path)
        print(f"Checado em = {today} em {len(ok)} matéria(s)")
    return errors


# ---------------------------------------------------------------- HTML -> planilha

def parse_html():
    h = open(HTML_PATH, encoding="utf-8").read()

    def block(name):
        m = block_re(name).search(h)
        if not m:
            sys.exit(f"Não encontrei o bloco MIDIA:{name} no index.html")
        return m.group(2)

    def g(rx, s):
        m = re.search(rx, s, re.S)
        return html.unescape(m.group(1)) if m else ""

    featured = {}
    for li in re.findall(r"<li>(.*?)</li>", block("DESTAQUES"), re.S):
        featured[g(r'href="([^"]*)"', li)] = (g(r'data-posicao="([^"]*)"', li), g(r"<small>(.*?)</small>", li))
    rows = []
    for li in re.findall(r"<li (.*?)</li>", block("LISTA"), re.S):
        url = g(r'href="([^"]*)"', li)
        c = g(r'data-cats="([^"]*)"', li).split()
        tg = [html.unescape(x) for x in re.findall(r"<span>(.*?)</span>", g(r'(<span class="tags">.*?</span></a>)', li))]
        titulo = g(r'class="ttl"[^>]*>(.*?)</span>', li)
        veiculo = g(r'class="src">(.*?)</span>', li)
        name = g(r'data-name="([^"]*)"', li)
        empresa = name[len(f"{veiculo} {titulo}"):].strip()
        r = {k: "" for k in KEYS}
        r.update({
            "publicar": "x", "data": g(r'data-date="([^"]*)"', li), "veiculo": veiculo, "titulo": titulo, "url": url,
            "categoria": c[0] if c else "", "formato": g(r'data-formato="([^"]*)"', li), "idioma": g(r'data-idioma="([^"]*)"', li),
            "assinado": "x" if "assinado" in c else "", "empresa": empresa, "paywall": "x" if TAG_PAYWALL in tg else "",
        })
        if url in featured:
            r["destaque"], r["nota"] = featured[url]
        rows.append(r)
    wordmarks = [html.unescape(x) for x in re.findall(r"<span>(.*?)</span>", g(r'class="wordmarks-list">(.*?)</span></a>', block("FAIXA") + "</a>"))]
    vehicles, order = {}, []
    for r in rows:
        if r["veiculo"] not in vehicles:
            vehicles[r["veiculo"]] = {"veiculo": r["veiculo"], "wordmark": "", "faixa": "", "site": "",
                                      "pais": "INT" if r["idioma"] == "en" else "BR"}
            order.append(r["veiculo"])
    for i, w in enumerate(wordmarks, 1):
        key = next((v for v in order if v.lower() == w.lower()), w)
        vehicles.setdefault(key, {"veiculo": key, "wordmark": "", "faixa": "", "site": "", "pais": ""})
        vehicles[key].update({"wordmark": w, "faixa": i})
    return rows, list(vehicles.values())


# ---------------------------------------------------------------- planilha

def style_sheet(ws, columns, center, limit=2000):
    head_fill = PatternFill("solid", fgColor="1F2937")
    for i, (key, header, width, note) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center" if key in center else "left", vertical="center", wrap_text=True)
        cell.comment = Comment(note, "midia.py")
        ws.column_dimensions[get_column_letter(i)].width = width
    keys = [c[0] for c in columns]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            if keys[c.column - 1] in center:
                c.alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 30
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(ws.max_row, 2)}"

    def listv(col_key, values, error):
        dv = DataValidation(type="list", formula1='"%s"' % ",".join(values), allow_blank=True, error=error, errorTitle="Valor inválido")
        ws.add_data_validation(dv)
        L = get_column_letter(keys.index(col_key) + 1)
        dv.add(f"{L}2:{L}{limit}")
    return listv


def write_xlsx(rows, vehicles, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append([c[1] for c in COLUMNS])
    for r in rows:
        ws.append([r.get(k, "") for k in KEYS])
    for L in (get_column_letter(KEYS.index(k) + 1) for k in ("data", "checado")):
        for c in ws[L][1:]:
            c.number_format = "@"  # data como texto, para o Excel não converter
    listv = style_sheet(ws, COLUMNS, FLAGS + ["destaque", "data", "categoria", "formato", "idioma", "checado"])
    for k in FLAGS:
        listv(k, ["x"], "Use x para marcar ou deixe vazio.")
    listv("destaque", POSICOES, f"Use um número de 1 a {DESTAQUES}.")
    listv("categoria", CAT_CODES, "Use perfil, tese, movimento ou portfolio.")
    listv("formato", list(FORMATOS), "Use texto, video ou audio.")
    listv("idioma", list(IDIOMAS), "Use pt ou en.")
    ws.freeze_panes = "F2"

    wv = wb.create_sheet(SHEET_V)
    wv.append([c[1] for c in COLUMNS_V])
    for v in vehicles:
        wv.append([v.get(k, "") for k in KEYS_V])
    listv_v = style_sheet(wv, COLUMNS_V, ["faixa", "pais"])
    listv_v("pais", sorted(PAISES), "Use BR ou INT.")
    wv.freeze_panes = "B2"
    # a coluna Veículo da aba Matérias só aceita o que está na aba Veículos
    dv = DataValidation(type="list", formula1=f"='{SHEET_V}'!$A$2:$A$500", allow_blank=True,
                        error="Cadastre o veículo na aba Veículos antes.", errorTitle="Veículo desconhecido")
    ws.add_data_validation(dv)
    L = get_column_letter(KEYS.index("veiculo") + 1)
    dv.add(f"{L}2:{L}2000")

    help_ws = wb.create_sheet("Como usar", 0)
    lines = [
        ("Na mídia: fonte da verdade da página /midia e da seção Na mídia da home", True),
        ("", False),
        ("Cada linha da aba Matérias é uma peça de imprensa. Edite aqui e depois peça ao Claude Code:", False),
        ("  \"Atualize o site a partir da planilha de mídia\"  (roda: python3 scripts/midia.py build e python3 scripts/pages.py build)", False),
        ("", False),
        ("Regra de entrada: você é nomeado na matéria E a peça é sobre você, assinada por você, ou um negócio da Headline que você liderou.", False),
        ("Ficam de fora (Publicar vazio): menção de passagem, lista de participantes, publieditorial e nota curta sem substância.", False),
        ("Publicar: x = aparece no site. Vazio = fica só na planilha, com o dossiê inteiro.", False),
        ("Destaque: 1 a 4 = os quatro cards da home, nessa ordem. Precisa haver exatamente quatro.", False),
        ("Categoria, pelo que a peça prova: perfil (quem você é), tese (como pensa), movimento (o que construiu), portfolio (o fundo operando).", False),
        ("  Formato, idioma e ano viram filtro na página, não categoria.", False),
        ("Assinado: x = texto seu. Vai para o bloco Textos assinados e entra no JSON-LD como author, não about.", False),
        ("subjectOf: x = entra no nó Person do JSON-LD, em todas as páginas. As cinco melhores, no máximo.", False),
        ("Título: como o veículo publicou. Datas em AAAA-MM-DD.", False),
        ("Aba Veículos: um veículo por linha. Faixa = posição na faixa de nomes da home; Wordmark = como o nome aparece nela.", False),
        ("Observações não vão para o site.", False),
        ("", False),
        ("Uma vez por trimestre: acrescente o que saiu, rode python3 scripts/midia.py check --links e revise os quatro destaques.", False),
        ("Passe o mouse no cabeçalho de cada coluna para ver a explicação. Não renomeie os cabeçalhos: o script procura as colunas pelo nome.", False),
    ]
    for text, bold in lines:
        help_ws.append([text])
        if bold:
            help_ws.cell(row=help_ws.max_row, column=1).font = Font(bold=True, size=14)
    help_ws.column_dimensions["A"].width = 130
    wb.active = wb.index(ws)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


# ---------------------------------------------------------------- CLI

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cmd", choices=["check", "build", "export"])
    ap.add_argument("--xlsx", default=XLSX_PATH)
    ap.add_argument("--force", action="store_true", help="export: sobrescreve a planilha existente")
    ap.add_argument("--links", action="store_true", help="check: abre cada URL publicada e atualiza Checado em")
    ap.add_argument("--todas", action="store_true", help="check --links: confere também as não publicadas")
    a = ap.parse_args()

    if a.cmd == "export":
        if os.path.exists(a.xlsx) and not a.force:
            sys.exit(f"{a.xlsx} já existe. Use --force para sobrescrever (edições feitas na planilha serão perdidas).")
        rows, vehicles = parse_html()
        write_xlsx(rows, vehicles, a.xlsx)
        print(f"{len(rows)} matérias e {len(vehicles)} veículos exportados para {os.path.relpath(a.xlsx, ROOT)}")
        return

    rows, vehicles = read_xlsx(a.xlsx)
    errors, warnings = validate(rows, vehicles)
    if a.links:
        errors += check_links(a.xlsx, rows, only_published=not a.todas)
    for w in warnings:
        print("aviso:", w)
    for e in errors:
        print("ERRO:", e)
    if errors:
        sys.exit(f"{len(errors)} erro(s). Nada foi alterado no index.html.")
    pub, changed = build(rows, vehicles, dry=(a.cmd == "check"))
    counts = " · ".join(f"{CAT_LABEL[c].lower()} {sum(r['categoria'] == c for r in pub)}" for c in CAT_CODES)
    print(f"{len(rows)} matérias na planilha · {len(pub)} publicadas ({counts}) · "
          f"{sum(r['assinado'] for r in pub)} assinadas · {len(strip_items(vehicles))} veículos na faixa")
    if a.cmd == "check":
        print("index.html " + ("seria alterado por build" if changed else "já está em dia com a planilha"))
    else:
        print("index.html atualizado" if changed else "index.html já estava em dia")


if __name__ == "__main__":
    main()
